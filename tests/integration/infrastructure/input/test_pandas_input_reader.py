"""Integration tests for real CSV/XLSX input files."""

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from csv_data_cleaner.domain.errors import InputDataError
from csv_data_cleaner.infrastructure.input import PandasInputReader
from csv_data_cleaner.infrastructure.mappers.data_frame_input_mapper import DataFrameInputMapper


def build_reader() -> PandasInputReader:
    return PandasInputReader(mapper=DataFrameInputMapper())


def active_sheet(workbook: Workbook) -> Worksheet:
    worksheet = workbook.active
    assert worksheet is not None
    return worksheet


def test_csv_and_xlsx_produce_same_canonical_rows(tmp_path: Path) -> None:
    csv_path = tmp_path / "customers.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(["name", "email"])
        writer.writerow(["Ada", "ada@example.com"])
        writer.writerow(["Grace", ""])

    xlsx_path = tmp_path / "customers.xlsx"
    workbook = Workbook()
    worksheet = active_sheet(workbook)
    worksheet.append(["name", "email"])
    worksheet.append(["Ada", "ada@example.com"])
    worksheet.append(["Grace", None])
    workbook.save(xlsx_path)

    reader = build_reader()

    assert reader.read(csv_path) == reader.read(xlsx_path)


def test_xlsx_reads_requested_sheet(tmp_path: Path) -> None:
    path = tmp_path / "customers.xlsx"
    workbook = Workbook()
    first = active_sheet(workbook)
    first.title = "First"
    first.append(["name"])
    first.append(["Wrong"])
    customers = workbook.create_sheet("Customers")
    customers.append(["name"])
    customers.append(["Ada"])
    workbook.save(path)

    result = build_reader().read(path, sheet="Customers")

    assert result.rows[0].values["name"] == "Ada"


@pytest.mark.parametrize("filename", ["missing.csv", "input.txt"])
def test_reader_rejects_unreadable_or_unsupported_input(tmp_path: Path, filename: str) -> None:
    path = tmp_path / filename
    if path.suffix == ".txt":
        path.write_text("name\nAda\n", encoding="utf-8")

    with pytest.raises(InputDataError):
        build_reader().read(path)
